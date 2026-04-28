import io
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse


def generate_books_pdf(books):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()
    elements = []

    # title
    title = Paragraph(
        'Library Management System - Books Report',
        styles['Title'],
    )
    elements.append(title)
    elements.append(Spacer(1, 20))

    # table headers
    headers = ['ID', 'Title', 'Author', 'Category', 'ISBN', 'Published Date', 'Copies', 'Available',]

    # table data
    data = [headers]
    for book in books:
        data.append([
            str(book.id),
            book.title,
            book.author.name,
            book.category.name if book.category else 'N/A',
            book.isbn,
            str(book.published_date),
            str(book.copies_available),
            'Yes' if book.copies_available > 0 else 'No',
        ])

    # create table
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        # header style
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2196F3')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),

        # data rows style
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),

        # alternating row colors
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [
            colors.white,
            colors.HexColor('#F5F5F5'),
        ]),

        # grid
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWHEIGHT', (0, 0), (-1, -1), 20),
    ]))

    elements.append(table)
    doc.build(elements)

    buffer.seek(0)
    return buffer


def generate_books_excel(books):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Books Report'

    # header style
    header_font = Font(
        bold=True,
        color='FFFFFF',
        size=11,
    )
    header_fill = PatternFill(
        start_color='2196F3',
        end_color='2196F3',
        fill_type='solid',
    )
    header_alignment = Alignment(
        horizontal='center',
        vertical='center',
    )
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    # headers
    headers = ['ID', 'Title', 'Author', 'Category', 'ISBN', 'Published Date', 'Copies Available', 'Is Available',]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border

    # data rows
    for row, book in enumerate(books, 2):
        data = [
            book.id,
            book.title,
            book.author.name,
            book.category.name if book.category else 'N/A',
            book.isbn,
            str(book.published_date),
            book.copies_available,
            'Yes' if book.copies_available > 0 else 'No',
        ]

        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(horizontal='center')
            cell.border = border

            # alternating row color
            if row % 2 == 0:
                cell.fill = PatternFill(
                    start_color='F5F5F5',
                    end_color='F5F5F5',
                    fill_type='solid',
                )

    # column widths
    column_widths = [5, 30, 20, 15, 15, 15, 10, 10]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[
            ws.cell(row=1, column=col).column_letter
        ].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer